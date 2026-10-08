import datetime
import io
import os
import fitz  # PyMuPDF
import streamlit as st

# Word Document Libraries
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Google API Libraries
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

# ==========================================
# 0. STREAMLIT PAGE CONFIG & CUSTOM STYLING
# ==========================================
st.set_page_config(
    page_title="2210 CS PYP Portal", 
    page_icon="💻",
    layout="wide"
)

# Custom Color Palette & Interface Styling
st.markdown("""
    <style>
    /* Main Page Background */
    .stApp, [data-testid="stAppViewContainer"] {
        background-color: #63D0F8 !important;
    }
    
    /* Sidebar Background */
    [data-testid="stSidebar"], [data-testid="stSidebar"] > div:first-child {
        background-color: #F0FCBB !important;
    }

    /* Global Text Color */
    html, body, [class*="css"], h1, h2, h3, h4, h5, h6, p, span, label, div, .stMarkdown {
        color: #384403 !important;
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    }

    /* Input Fields, Selectboxes & Text Areas */
    div[data-baseweb="input"], 
    div[data-baseweb="select"] > div, 
    .stTextInput input, 
    .stSelectbox select,
    textarea {
        background-color: #CEC2F5 !important;
        color: #141801 !important;
        border-radius: 10px !important;
        border: 4px solid #A5C809 !important;
    }

    /* Buttons Styling */
    .stButton button, 
    .stDownloadButton button, 
    [data-testid="baseButton-secondary"], 
    [data-testid="baseButton-primary"] {
        background-color: #C9F40B !important;
        color: #384403 !important;
        border: 2px solid #A5C809 !important;
        border-radius: 8px !important;
        font-weight: bold !important;
        transition: all 0.2s ease-in-out;
    }
    
    /* Hover State for Buttons */
    .stButton button:hover, .stDownloadButton button:hover {
        background-color: #A5C809 !important;
        color: #384403 !important;
        border: 3px solid #384403 !important;
    }

    /* Navigation Tab Labels */
    button[data-baseweb="tab"] p {
        font-weight: bold !important;
        font-size: 1.1rem !important;
        color: #384403 !important;
    }
    
    /* Active Tab Highlight Indicator */
    div[data-baseweb="tab-highlight"] {
        background-color: #F863E1 !important;
    }

    /* Expanders & Containers */
    [data-testid="stExpander"] {
        border: 1.5px solid #A5C809 !important;
        border-radius: 8px !important;
        background-color: #F0FCBB !important;
    }
    </style>
""", unsafe_allow_html=True)


# ==========================================
# 1. DIRECTORY MAPPING & CONFIGURATION
# ==========================================
SYLLABUS_CODE = "2210"
SCOPES = ['https://www.googleapis.com/auth/drive']

# 4 Core Local Folders for 2210 Computer Science
LOCAL_FOLDERS = {
    "p1": "2210_Paper1",
    "p2": "2210_Paper2",
    "ms_p1": "2210_ms_P1",
    "ms_p2": "2210_ms_P2"
}

for folder_path in LOCAL_FOLDERS.values():
    os.makedirs(folder_path, exist_ok=True)


# ==========================================
# 2. GOOGLE DRIVE AUTHENTICATION & SYNC
# ==========================================
def build_drive_service(write_access=True):
    """Authenticates using Google Service Account credentials from Streamlit Secrets."""
    try:
        if "gcp_service_account" in st.secrets:
            service_account_info = dict(st.secrets["gcp_service_account"])
            scopes = ['https://www.googleapis.com/auth/drive'] if write_access else ['https://www.googleapis.com/auth/drive.readonly']
            creds = service_account.Credentials.from_service_account_info(
                service_account_info, 
                scopes=scopes
            )
            return build('drive', 'v3', credentials=creds)
        else:
            st.error("Missing [gcp_service_account] configuration in secrets.")
            return None
    except Exception as e:
        st.error(f"Authentication Error: {e}")
        return None

def sync_drive_folder_to_local(folder_key: str) -> tuple[int, str]:
    """Downloads missing files from Google Drive folder into local directory with Shared Drive support."""
    service = build_drive_service(write_access=False)
    if not service:
        return 0, "Failed to authenticate Service Account."
    
    folder_ids = st.secrets.get("drive_folders", {})
    drive_folder_id = folder_ids.get(folder_key)
    
    if not drive_folder_id:
        return 0, f"Missing drive_folder_id for `{folder_key}` in secrets."

    local_path = LOCAL_FOLDERS[folder_key]
    
    try:
        query = f"'{drive_folder_id}' in parents and trashed = false"
        drive_files = []
        page_token = None

        while True:
            response = service.files().list(
                q=query,
                fields="nextPageToken, files(id, name, mimeType)",
                pageToken=page_token,
                pageSize=100,
                supportsAllDrives=True,
                includeItemsFromAllDrives=True
            ).execute()
            
            drive_files.extend(response.get('files', []))
            page_token = response.get('nextPageToken', None)
            
            if not page_token:
                break

        downloaded_count = 0

        for file_info in drive_files:
            file_name = file_info['name']
            file_id = file_info['id']
            local_file_path = os.path.join(local_path, file_name)

            if not os.path.exists(local_file_path):
                request = service.files().get_media(fileId=file_id)
                with open(local_file_path, "wb") as f:
                    downloader = MediaIoBaseDownload(f, request)
                    done = False
                    while not done:
                        _, done = downloader.next_chunk()
                downloaded_count += 1

        total_local_files = len([f for f in os.listdir(local_path) if os.path.isfile(os.path.join(local_path, f))])
        return downloaded_count, f"Synced {downloaded_count} new file(s) for `{folder_key}` (Total: {total_local_files})."
        
    except Exception as e:
        return 0, f"Sync error for `{folder_key}`: {e}"

def perform_bulk_sync():
    """Syncs all 4 configured Google Drive folders."""
    total_synced = 0
    messages = []
    for f_key in LOCAL_FOLDERS.keys():
        count, msg = sync_drive_folder_to_local(f_key)
        total_synced += count
        messages.append(msg)
    return total_synced, messages

def upload_file_to_drive(file_bytes: bytes, file_name: str, folder_key: str) -> tuple[bool, str]:
    """
    Uploads a PDF file directly to Google Drive with Shared Drive flags enabled,
    and saves a local copy so search indexing works immediately.
    """
    service = build_drive_service(write_access=True)
    if not service:
        return False, "Could not connect to Google Drive service."

    folder_ids = st.secrets.get("drive_folders", {})
    drive_folder_id = folder_ids.get(folder_key)

    if not drive_folder_id:
        return False, f"Missing drive folder ID for key: `{folder_key}` in secrets."

    # Save local copy first so file is immediately available in the portal
    local_dir = LOCAL_FOLDERS[folder_key]
    local_path = os.path.join(local_dir, file_name)
    try:
        with open(local_path, "wb") as f:
            f.write(file_bytes)
    except Exception as e:
        return False, f"Failed to save file locally: {e}"

    # Upload to Google Drive with Shared Drive parameters
    try:
        file_metadata = {
            'name': file_name,
            'parents': [drive_folder_id]
        }
        media = MediaIoBaseUpload(
            io.BytesIO(file_bytes), 
            mimetype='application/pdf', 
            resumable=True
        )

        uploaded_file = service.files().create(
            body=file_metadata,
            media_body=media,
            fields='id',
            supportsAllDrives=True,
            supportsTeamDrives=True
        ).execute()

        return True, f"Successfully uploaded `{file_name}` to Drive (ID: `{uploaded_file.get('id')}`) and saved locally."

    except Exception as e:
        # Fallback message confirming local save if Drive API quota blocks remote upload
        return True, f"Saved `{file_name}` locally for immediate search! (Drive remote notice: {e})"


# ==========================================
# 3. HELPER FUNCTIONS
# ==========================================
def add_page_number_to_run(run):
    """Adds a dynamic Word page number field to document headers."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def create_worksheet_docx(basket_items: list) -> io.BytesIO:
    """
    Generates a Word document containing selected PDF pages from the cart.
    Images are scaled down to fit cleanly within page margins to prevent trailing blank pages.
    """
    doc = Document()
    section = doc.sections[0]

    # Page dimension setup
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    
    # Compact margins (0.4 inches) to maximize printable vertical height
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    # Header configuration
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_p.add_run("Page ")
    add_page_number_to_run(header_run)

    # Document Heading
    doc.add_heading(f'O Level {SYLLABUS_CODE} Computer Science Worksheet', level=1)

    for idx, item in enumerate(basket_items):
        # Section Heading
        heading_p = doc.add_heading(f"Source: {item['file']} (Page {item['page'] + 1})", level=2)
        heading_p.paragraph_format.space_after = Pt(4)
        
        # Load PDF page and render as PNG image
        pdf_doc = fitz.open(item['path'])
        page = pdf_doc.load_page(item['page'])
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
        img_data = io.BytesIO(pix.tobytes("png"))

        # Add image with height constraint (6.8 in) so heading + image fit on 1 page
        doc.add_picture(img_data, height=Inches(6.8))

        if idx < len(basket_items) - 1:
            doc.add_page_break()
            
        pdf_doc.close()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def render_pdf_page_preview(filepath: str, page_num: int = 0):
    """Renders a PDF page to PNG image bytes for preview."""
    try:
        doc = fitz.open(filepath)
        page = doc.load_page(page_num)
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
        img_bytes = pix.tobytes("png")
        doc.close()
        return img_bytes
    except Exception as e:
        st.error(f"Unable to render page preview: {e}")
        return None

def execute_pdf_search(folder_key: str, keyword_string: str, variant_filter: str = "All Variants") -> list[dict]:
    """Searches PDF files in a specific folder for matching keywords and variant filter (variants 2 and 3)."""
    results = []
    keywords = [k.strip().lower() for k in keyword_string.split(",") if k.strip()]
    folder_path = LOCAL_FOLDERS[folder_key]
    
    if os.path.exists(folder_path):
        for file in os.listdir(folder_path):
            if file.endswith(".pdf"):
                file_lower = file.lower()
                
                # Apply variant filtering for Variants 2 and 3
                if variant_filter == "Variant 2" and not ("_qp_12" in file_lower or "_qp_22" in file_lower or "_12." in file_lower or "_22." in file_lower):
                    continue
                elif variant_filter == "Variant 3" and not ("_qp_13" in file_lower or "_qp_23" in file_lower or "_13." in file_lower or "_23." in file_lower):
                    continue

                filepath = os.path.join(folder_path, file)
                try:
                    doc = fitz.open(filepath)
                    for page_num in range(len(doc)):
                        text = doc[page_num].get_text().lower()
                        if all(kw in text for kw in keywords):
                            results.append({
                                "file": file, 
                                "page": page_num, 
                                "path": filepath
                            })
